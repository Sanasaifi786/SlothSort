from pathlib import Path
import docx
from pypdf import PdfWriter
from slothsort.extract import extract_text_snippet


def test_extract_plaintext(tmp_path: Path):
    txt_file = tmp_path / "sample.txt"
    txt_file.write_text("Hello SlothSort! This is a simple test document.", encoding="utf-8")

    snippet = extract_text_snippet(txt_file, max_chars=100)
    assert "Hello SlothSort!" in snippet


def test_extract_docx(tmp_path: Path):
    doc_file = tmp_path / "resume.docx"
    doc = docx.Document()
    doc.add_heading("Senior Python Developer", level=1)
    doc.add_paragraph("Experienced software engineer specializing in offline AI agents.")
    doc.save(str(doc_file))

    snippet = extract_text_snippet(doc_file, max_chars=100)
    assert "Senior Python Developer" in snippet
    assert "Experienced software engineer" in snippet


def test_extract_pdf(tmp_path: Path):
    pdf_file = tmp_path / "invoice.pdf"
    writer = PdfWriter()
    # Add a blank page with no crash
    writer.add_blank_page(width=72, height=72)
    with open(pdf_file, "wb") as f:
        writer.write(f)

    snippet = extract_text_snippet(pdf_file)
    assert isinstance(snippet, str)


def test_extract_corrupted_file(tmp_path: Path):
    fake_pdf = tmp_path / "broken.pdf"
    fake_pdf.write_bytes(b"NOT A REAL PDF CONTENT AT ALL")

    snippet = extract_text_snippet(fake_pdf)
    assert snippet == ""


def test_extract_nonexistent_file(tmp_path: Path):
    missing = tmp_path / "ghost.pdf"
    snippet = extract_text_snippet(missing)
    assert snippet == ""


def test_extract_max_chars_limit(tmp_path: Path):
    long_txt = tmp_path / "long.txt"
    long_txt.write_text("word " * 500, encoding="utf-8")

    snippet = extract_text_snippet(long_txt, max_chars=50)
    assert len(snippet) <= 50
