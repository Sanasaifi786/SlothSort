"""Script to generate a realistic test_downloads folder with ~40 messy files

and a ground-truth labels.csv file for accuracy benchmarking.
"""

import csv
import os
import shutil
import time
from pathlib import Path
import docx
from pypdf import PdfWriter


def create_mock_pdf(path: Path, content_text: str) -> None:
    """Create a minimal PDF file with metadata or text content."""
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.add_metadata({"/Title": content_text[:100], "/Subject": content_text[:200]})
    with open(path, "wb") as f:
        writer.write(f)


def create_mock_docx(path: Path, heading: str, body: str) -> None:
    """Create a minimal DOCX file with headings and text."""
    doc = docx.Document()
    doc.add_heading(heading, level=1)
    doc.add_paragraph(body)
    doc.save(str(path))


def generate_test_files(target_dir: Path) -> list[tuple[str, str]]:
    """Build ~40 messy files representing a typical chaotic Downloads folder."""
    target_dir.mkdir(parents=True, exist_ok=True)

    items: list[tuple[str, str, str, str]] = [
        # (filename, category, type, content/heading)
        # --- PHOTOS (Rules: .jpg, .png, .webp) ---
        ("IMG_2231.JPG", "Photos", "bin", "fake jpeg bytes"),
        ("Screenshot_2026-10-22_141102.png", "Photos", "bin", "fake png bytes"),
        ("meme_final(1).webp", "Photos", "bin", "fake webp bytes"),
        ("DSC_0042.jpeg", "Photos", "bin", "camera photo"),
        ("holiday_sunset.jpg", "Photos", "bin", "sunset photo"),
        ("scan_receipt_photo.png", "Photos", "bin", "photo receipt"),
        
        # --- INSTALLERS (Rules: .exe, .msi) ---
        ("vscode_setup_x64.exe", "Installers", "bin", "installer binary"),
        ("git_installer_v2.55.exe", "Installers", "bin", "git setup"),
        ("python_3.12_amd64.exe", "Installers", "bin", "python installer"),
        ("docker_desktop_setup.msi", "Installers", "bin", "docker msi"),
        ("zoom_client_installer.msi", "Installers", "bin", "zoom client"),

        # --- ARCHIVES (Rules: .zip, .7z, .tar.gz) ---
        ("project_backup_oct2026.zip", "Archives", "bin", "zip archive"),
        ("dataset_photos_raw.7z", "Archives", "bin", "7z archive"),
        ("linux_kernel_sources.tar.gz", "Archives", "bin", "tarball"),
        ("icons_bundle(2).zip", "Archives", "bin", "zip icons"),

        # --- BILLS (Ambiguous: PDF / TXT -> AI needed) ---
        ("Airtel_bill_oct(3).pdf", "Bills", "txt", "Airtel Telecommunications Invoice Monthly Broadband Bill Rs 999 Due Date 2026-11-05 Account 492019"),
        ("bescom_electricity_receipt.pdf", "Bills", "txt", "Electricity Board Bangalore BESCOM Consumer Bill Payment Receipt Meter Number 882199 Amount Due Rs 1450"),
        ("amazon_tax_invoice_INV2026.txt", "Bills", "txt", "Amazon Retail India Tax Invoice Order ID 402-991283 Sold by Appario Retail Total INR 3299"),
        ("credit_card_hdfc_statement.pdf", "Bills", "txt", "HDFC Bank Credit Card Account Statement Minimum Due Rs 4500 Total Amount Outstanding Rs 32410"),
        ("water_utility_october.pdf", "Bills", "txt", "Municipal Water Supply Corporation Monthly Utility Consumption Bill Charges Rs 480"),

        # --- STUDY (Ambiguous: PDF / DOCX / TXT -> AI needed) ---
        ("DS_unit3_notes_final.pdf", "Study", "txt", "Data Structures and Algorithms Unit 3 Lecture Notes Binary Search Trees AVL Trees Graphs BFS DFS traversal"),
        ("operating_systems_cheatsheet.docx", "Study", "docx", "Operating Systems Virtual Memory Paging Deadlock Prevention Mutex Semaphore scheduling algorithms"),
        ("compiler_design_lab4.txt", "Study", "txt", "Compiler Design Lab Assignment 4 Lexical Analyzer Yacc Lex Grammar parsing shift-reduce parser"),
        ("linear_algebra_eigenvalues.pdf", "Study", "txt", "Mathematics Department Linear Algebra Eigenvalues and Eigenvectors Characteristic polynomial diagonalization"),
        ("machine_learning_summary_notes.docx", "Study", "docx", "Machine Learning Supervised Learning Cost Function Gradient Descent Overfitting Regularization"),

        # --- CAREER (Ambiguous: PDF / DOCX -> AI needed) ---
        ("resume_v7_updated_final.docx", "Career", "docx", "Professional Software Engineer Resume 4 years experience Python Docker React Distributed Systems"),
        ("cover_letter_google.docx", "Career", "docx", "Cover Letter Dear Hiring Manager I am writing to express my strong interest in the Software Engineer role at Google"),
        ("offer_letter_techcorp_confidential.pdf", "Career", "txt", "Employment Offer Letter TechCorp Technologies Candidate Salary CTC Joining Date 2026-12-01"),
        ("experience_letter_signed.pdf", "Career", "txt", "Certificate of Experience This is to certify that candidate worked as Software Developer from 2024 to 2026"),
        ("cv_sana_saifi_2026.pdf", "Career", "txt", "Curriculum Vitae Sana Saifi Computer Science Specialist Full Stack Development Open Source Contributor"),

        # --- DOCUMENTS & OTHER ---
        ("meeting_minutes_team_sync.txt", "Documents", "txt", "Weekly Engineering Sync Meeting Minutes Sprint review action items assigned to frontend backend team"),
        ("apartment_lease_agreement.pdf", "Documents", "txt", "Residential Rental Lease Agreement Tenant and Landlord Terms and Conditions Deposit Security"),
        ("passport_application_acknowledgement.pdf", "Documents", "txt", "Passport Seva Kendra Appointment Confirmation and Document Verification Application Receipt"),
        ("server_access_instructions.txt", "Documents", "txt", "Server SSH Configuration port forwarding bastion host credentials instructions guide"),
        ("random_scratchpad_ideas.txt", "Other", "txt", "Just some random notes buy groceries milk bread check gym timings call car mechanic"),
    ]

    labels: list[tuple[str, str]] = []

    # Touch each file with an mtime 2 hours in the past so recent-check allows processing
    old_timestamp = time.time() - 7200

    for filename, category, file_type, text_content in items:
        file_path = target_dir / filename
        if file_type == "bin":
            file_path.write_bytes(text_content.encode("utf-8"))
        elif file_type == "docx":
            create_mock_docx(file_path, heading=filename, body=text_content)
        elif file_type == "txt":
            # For testing convenience and fast accurate text extraction:
            file_path.write_text(text_content, encoding="utf-8")

        # Set modified time in past
        os.utime(file_path, (old_timestamp, old_timestamp))
        labels.append((filename, category))

    # Write labels.csv
    labels_csv_path = target_dir / "labels.csv"
    with open(labels_csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["filename", "expected_category"])
        writer.writerows(labels)

    return labels


def main():
    target = Path("test_downloads").resolve()
    if target.exists():
        shutil.rmtree(target)
    labels = generate_test_files(target)
    print(f"Generated {len(labels)} test files in '{target}'.")
    print(f"Created ground-truth labels at '{target / 'labels.csv'}'.")


if __name__ == "__main__":
    main()
